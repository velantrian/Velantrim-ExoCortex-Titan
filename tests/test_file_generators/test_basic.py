"""
tests/test_file_generators/test_basic.py
==========================================
Базовые тесты всех генераторов:
- Создание не падает
- Файлы реально создаются
- Файлы не пустые
- GenerationResult корректный
- Темы применяются

Зависимости опциональные — тесты используют pytest.importorskip
для пропуска при отсутствии библиотек.
"""

import os

import pytest

# ─── Fixtures ─────────────────────────────────────────────────────────────────

@pytest.fixture
def simple_spec():
    """Минимальная GenerationSpec со всеми типами блоков."""
    from core.file_generators import (
        CalloutBlock,
        CodeBlock,
        DividerBlock,
        DocumentMetadata,
        FactBlock,
        GenerationSpec,
        HeadingBlock,
        ListBlock,
        ParagraphBlock,
        QuoteBlock,
        TableBlock,
    )

    return GenerationSpec(
        metadata=DocumentMetadata(
            title="Test Document",
            author="Test Suite",
            subject="Unit testing",
            keywords=["test", "velantrim"],
        ),
        theme="velantrim",
        blocks=[
            HeadingBlock(text="Раздел 1", level=1),
            ParagraphBlock(text="Это обычный параграф."),
            ParagraphBlock(text="А это жирный.", style="bold"),
            ParagraphBlock(text="А это курсив.", style="italic"),
            CalloutBlock(
                callout_type="info",
                title="Внимание",
                text="Это callout-блок.",
            ),
            HeadingBlock(text="Список", level=2),
            ListBlock(items=["Первый", "Второй", "Третий"]),
            ListBlock(items=["Раз", "Два"], ordered=True),
            HeadingBlock(text="Таблица", level=2),
            TableBlock(
                headers=["Имя", "Возраст", "Город"],
                rows=[
                    ["Алиса", 30, "Москва"],
                    ["Боб", 25, "Лондон"],
                ],
                caption="Тестовая таблица",
            ),
            HeadingBlock(text="Код", level=2),
            CodeBlock(
                code="def hello():\n    return 'Velantrim'",
                language="python",
                caption="Пример функции",
            ),
            HeadingBlock(text="Цитата", level=2),
            QuoteBlock(
                text="Память без проверки — это просто хранение.",
                author="Velantrim Manifest",
            ),
            DividerBlock(),
            HeadingBlock(text="🔱 Факты", level=2),
            FactBlock(
                fact_id="test_f1",
                claim="Velantrim работает с памятью AI агентов",
                confidence=0.99,
                epistemic_state="Validated",
                source="test_suite",
            ),
            FactBlock(
                fact_id="test_f2",
                claim="Гипотеза которая пока не подтверждена",
                confidence=0.55,
                epistemic_state="Hypothesized",
                source="speculation",
            ),
        ],
    )


@pytest.fixture
def facts_list():
    """Список фактов для тестирования export_facts."""
    return [
        {
            "fact_id": f"fact_{i}",
            "claim": f"Утверждение номер {i}",
            "confidence": 0.5 + i * 0.05,
            "epistemic_state": "Validated" if i % 2 == 0 else "Hypothesized",
            "source": f"source_{i % 3}",
        }
        for i in range(10)
    ]


# ═══════════════════════════════════════════════════════════════════════════════
# 1. PDF Generator
# ═══════════════════════════════════════════════════════════════════════════════

class TestPDFGenerator:
    """Тесты для PDFGenerator."""

    def test_import(self):
        from core.file_generators import FileExporter
        assert FileExporter is not None

    def test_pdf_generation(self, simple_spec, tmp_path):
        """PDF файл создаётся и не пустой."""
        pytest.importorskip("reportlab")
        from core.file_generators import FileExporter

        output = str(tmp_path / "test.pdf")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None, f"Error: {result.error}"
        assert result.format == "pdf"
        assert os.path.exists(output)
        assert os.path.getsize(output) > 1000   # PDF не пустой
        assert result.file_size_bytes > 1000

    def test_pdf_with_all_themes(self, simple_spec, tmp_path):
        """Все 5 тем рендерятся без ошибок."""
        pytest.importorskip("reportlab")
        from core.file_generators import FileExporter

        for theme in ["clean", "scientific", "business", "dark", "velantrim"]:
            simple_spec.theme = theme
            output = str(tmp_path / f"test_{theme}.pdf")
            result = FileExporter().export(simple_spec, output)
            assert result.error is None, f"Theme {theme} failed: {result.error}"
            assert os.path.getsize(output) > 1000

    def test_pdf_from_facts(self, facts_list, tmp_path):
        """Shortcut export_facts работает."""
        pytest.importorskip("reportlab")
        from core.file_generators import FileExporter

        output = str(tmp_path / "facts.pdf")
        result = FileExporter().export_facts(
            facts_list, output, title="Test Facts Report",
        )
        assert result.error is None
        assert os.path.getsize(output) > 1000


# ═══════════════════════════════════════════════════════════════════════════════
# 2. DOCX Generator
# ═══════════════════════════════════════════════════════════════════════════════

class TestDOCXGenerator:

    def test_docx_generation(self, simple_spec, tmp_path):
        pytest.importorskip("docx")
        from core.file_generators import FileExporter

        output = str(tmp_path / "test.docx")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None, f"Error: {result.error}"
        assert result.format == "docx"
        assert os.path.exists(output)
        assert os.path.getsize(output) > 1000

    def test_docx_round_trip(self, simple_spec, tmp_path):
        """Создаём DOCX и читаем обратно — проверяем что python-docx может его открыть."""
        pytest.importorskip("docx")
        import docx as python_docx

        from core.file_generators import FileExporter

        output = str(tmp_path / "round.docx")
        FileExporter().export(simple_spec, output)

        # Должен открыться без ошибок
        doc = python_docx.Document(output)
        # Должны быть какие-то параграфы
        assert len(doc.paragraphs) > 0


# ═══════════════════════════════════════════════════════════════════════════════
# 3. PPTX Generator
# ═══════════════════════════════════════════════════════════════════════════════

class TestPPTXGenerator:

    def test_pptx_generation(self, simple_spec, tmp_path):
        pytest.importorskip("pptx")
        from core.file_generators import FileExporter

        output = str(tmp_path / "test.pptx")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None
        assert result.format == "pptx"
        assert os.path.exists(output)
        assert os.path.getsize(output) > 5000   # PPTX крупнее PDF
        # Должно быть несколько слайдов (минимум title + контент)
        assert result.page_count is None or result.page_count >= 1

    def test_pptx_slide_dimensions(self, simple_spec, tmp_path):
        """Слайды 16:9 widescreen."""
        pytest.importorskip("pptx")
        from pptx import Presentation

        from core.file_generators import FileExporter

        output = str(tmp_path / "wide.pptx")
        FileExporter().export(simple_spec, output)

        prs = Presentation(output)
        # 13.333 inches = 12192000 EMU
        assert prs.slide_width > 12_000_000


# ═══════════════════════════════════════════════════════════════════════════════
# 4. XLSX Generator
# ═══════════════════════════════════════════════════════════════════════════════

class TestXLSXGenerator:

    def test_xlsx_generation(self, simple_spec, tmp_path):
        pytest.importorskip("openpyxl")
        from core.file_generators import FileExporter

        output = str(tmp_path / "test.xlsx")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None
        assert result.format == "xlsx"
        assert os.path.exists(output)
        assert os.path.getsize(output) > 1000

    def test_xlsx_multiple_sheets(self, facts_list, tmp_path):
        """Из фактов создаются несколько листов."""
        pytest.importorskip("openpyxl")
        from openpyxl import load_workbook

        from core.file_generators import FileExporter

        output = str(tmp_path / "multi.xlsx")
        FileExporter().export_facts(facts_list, output)

        wb = load_workbook(output)
        # Должны быть Summary, Facts (минимум)
        sheet_names = [s.title for s in wb.worksheets]
        assert any("Summary" in name for name in sheet_names)
        assert any("Facts" in name for name in sheet_names)


# ═══════════════════════════════════════════════════════════════════════════════
# 5. HTML Generator — без зависимостей, всегда должен работать
# ═══════════════════════════════════════════════════════════════════════════════

class TestHTMLGenerator:

    def test_html_generation(self, simple_spec, tmp_path):
        """HTML работает без внешних зависимостей."""
        from core.file_generators import FileExporter

        output = str(tmp_path / "test.html")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None
        assert result.format == "html"
        assert os.path.exists(output)

        # Проверим что это валидный HTML
        with open(output, encoding="utf-8") as f:
            content = f.read()
        assert "<!DOCTYPE html>" in content
        assert "<html" in content
        assert "</html>" in content
        assert simple_spec.metadata.title in content

    def test_html_self_contained(self, simple_spec, tmp_path):
        """HTML standalone — нет ссылок на внешние ресурсы."""
        from core.file_generators import FileExporter
        output = str(tmp_path / "standalone.html")
        FileExporter().export(simple_spec, output)

        with open(output, encoding="utf-8") as f:
            content = f.read()
        # CSS inline
        assert "<style>" in content
        # Нет CDN
        assert "cdn." not in content.lower()
        assert "googleapis.com" not in content
        assert "cloudflare" not in content.lower()

    def test_html_themes(self, simple_spec, tmp_path):
        """Все темы применяются через CSS custom properties."""
        from core.file_generators import FileExporter

        for theme in ["clean", "scientific", "business", "dark", "velantrim"]:
            simple_spec.theme = theme
            output = str(tmp_path / f"theme_{theme}.html")
            FileExporter().export(simple_spec, output)
            with open(output, encoding="utf-8") as f:
                content = f.read()
            assert "--color-primary" in content


# ═══════════════════════════════════════════════════════════════════════════════
# 6. Markdown Generator — без зависимостей
# ═══════════════════════════════════════════════════════════════════════════════

class TestMarkdownGenerator:

    def test_markdown_generation(self, simple_spec, tmp_path):
        from core.file_generators import FileExporter
        output = str(tmp_path / "test.md")
        result = FileExporter().export(simple_spec, output)

        assert result.error is None
        assert os.path.exists(output)

        with open(output, encoding="utf-8") as f:
            content = f.read()

        # YAML frontmatter
        assert content.startswith("---")
        # Содержит заголовок
        assert simple_spec.metadata.title in content
        # Markdown heading syntax
        assert "# " in content
        # Markdown table syntax
        assert "|" in content


# ═══════════════════════════════════════════════════════════════════════════════
# 7. FileExporter — главный оркестратор
# ═══════════════════════════════════════════════════════════════════════════════

class TestFileExporter:

    def test_supported_formats(self):
        from core.file_generators import FileExporter
        formats = FileExporter().get_supported_formats()
        # Минимум: html, md, markdown (без зависимостей)
        assert ".html" in formats
        assert ".md" in formats

    def test_unsupported_extension(self, simple_spec, tmp_path):
        """Несуществующий формат → понятная ошибка."""
        from core.file_generators import FileExporter
        output = str(tmp_path / "test.xyz")
        result = FileExporter().export(simple_spec, output)
        assert result.error is not None
        assert ".xyz" in result.error

    def test_export_multi(self, simple_spec, tmp_path):
        """Несколько форматов одним вызовом."""
        from core.file_generators import FileExporter
        base = str(tmp_path / "multi")
        # Только форматы без зависимостей
        results = FileExporter().export_multi(
            simple_spec, base, formats=["html", "md"],
        )
        assert "html" in results
        assert "md" in results
        for fmt, result in results.items():
            assert result.error is None, f"Format {fmt}: {result.error}"
            assert os.path.exists(result.output_path)

    def test_stats(self, simple_spec, tmp_path):
        """Статистика обновляется после генерации."""
        from core.file_generators import FileExporter
        exporter = FileExporter()
        exporter.export(simple_spec, str(tmp_path / "stats1.html"))
        exporter.export(simple_spec, str(tmp_path / "stats2.md"))

        stats = exporter.get_stats()
        assert stats["generated"] >= 2
        assert stats["total_bytes"] > 0
        # На быстрых машинах время может округлиться до 0 ms
        assert stats["total_time_ms"] >= 0


# ═══════════════════════════════════════════════════════════════════════════════
# 8. Базовые модели данных
# ═══════════════════════════════════════════════════════════════════════════════

class TestDataModels:

    def test_generation_spec_from_facts(self, facts_list):
        from core.file_generators import GenerationSpec
        spec = GenerationSpec.from_facts(
            facts_list, title="Test", theme="velantrim",
        )
        assert spec.metadata.title == "Test"
        assert spec.theme == "velantrim"
        # Должны быть FactBlock'и
        from core.file_generators import FactBlock, HeadingBlock
        fact_blocks = [b for b in spec.blocks if isinstance(b, FactBlock)]
        assert len(fact_blocks) == len(facts_list)
        # И заголовок
        assert any(isinstance(b, HeadingBlock) for b in spec.blocks)

    def test_themes_available(self):
        from core.file_generators import THEMES, get_theme
        expected = {"clean", "scientific", "business", "dark", "velantrim"}
        assert expected.issubset(set(THEMES.keys()))

        # get_theme с неизвестным именем → fallback на clean
        theme = get_theme("nonexistent")
        assert theme.name == "clean"

    def test_blocks_dataclass(self):
        """Все блоки — это dataclass'ы с правильными default."""
        from core.file_generators import (
            CalloutBlock,
            CodeBlock,
            DividerBlock,
            FactBlock,
            HeadingBlock,
            ImageBlock,
            ListBlock,
            ParagraphBlock,
            QuoteBlock,
            TableBlock,
        )

        # Каждый блок должен иметь block_type
        assert HeadingBlock(text="x").block_type == "heading"
        assert ParagraphBlock(text="x").block_type == "paragraph"
        assert ListBlock(items=["x"]).block_type == "list"
        assert TableBlock(headers=["a"], rows=[["b"]]).block_type == "table"
        assert CodeBlock(code="x").block_type == "code"
        assert ImageBlock(path="x.png").block_type == "image"
        assert CalloutBlock(text="x").block_type == "callout"
        assert QuoteBlock(text="x").block_type == "quote"
        assert DividerBlock().block_type == "divider"
        assert FactBlock(fact_id="x", claim="y").block_type == "fact"


# ═══════════════════════════════════════════════════════════════════════════════
# 9. Velantrim Reports (готовые шаблоны)
# ═══════════════════════════════════════════════════════════════════════════════

class TestVelantrimReports:

    def test_knowledge_base_generation(self, facts_list, tmp_path):
        """generate_knowledge_base создаёт валидный spec."""
        from core.file_generators import FileExporter, GenerationSpec
        from core.velantrim_reports import generate_knowledge_base

        spec = generate_knowledge_base(facts_list, theme="scientific")
        assert isinstance(spec, GenerationSpec)
        assert spec.theme == "scientific"
        # Должен быть HTML экспорт
        output = str(tmp_path / "kb.html")
        result = FileExporter().export(spec, output)
        assert result.error is None
        assert os.path.exists(output)

    def test_sprint_review_generation(self, tmp_path):
        from core.file_generators import FileExporter
        from core.velantrim_reports import generate_sprint_review

        sprint = {
            "number": "test",
            "name": "Test Sprint",
            "team": "Test Team",
            "goal": "Test the test",
            "delivered": [{"title": "Test feature", "type": "feature"}],
            "metrics": [["Tests", 100, 200, "+100%"]],
            "carryover": ["Carry-over item"],
            "next_goals": ["Next goal"],
        }
        spec = generate_sprint_review(sprint, format_hint="docx")
        output = str(tmp_path / "sprint.html")
        result = FileExporter().export(spec, output)
        assert result.error is None
